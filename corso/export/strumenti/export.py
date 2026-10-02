"""Esporta un deck (cartella con project/deck.json e project/slides/*.html) in HTML, PDF e PPTX.
uso: python3 export.py <cartella deck> <nome file> <cartella uscita>"""
import sys, os, re, json, html as H, io
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu
from PIL import Image

deck, name, out = sys.argv[1], sys.argv[2], sys.argv[3]
import glob
CHROME = (glob.glob('/opt/pw-browsers/chromium-*/chrome-linux/chrome') or [None])[0]
d = json.load(open(f'{deck}/project/deck.json'))
title = d['title']
links = ''.join(f'<link rel="stylesheet" href="{v["href"]}">' for v in d.get('faces', {}).values() if 'href' in v)
BASE = ('*{margin:0;box-sizing:border-box} section.s > *:not(svg):not([style*="position:absolute"]){position:relative} '
        'section.s [style*="margin"]{margin:0 !important} section.s ul,section.s ol{padding-left:40px} '
        '')

slides, notes = [], []
for i in d['order']:
    h = open(f'{deck}/project/slides/{i}.html').read()
    m = re.search(r'<aside>(.*?)</aside>', h, re.S)
    notes.append(H.unescape(re.sub(r'<[^>]+>', '', m.group(1))).strip() if m else '')
    h = re.sub(r'<aside>.*?</aside>', '', h, flags=re.S)
    h = h.replace('<section ', '<section class="s" ', 1).replace('style="', 'style="position:relative; width:1920px; height:1080px; overflow:hidden; ', 1)
    slides.append(h)

# ---------- HTML da consultare (scala alla finestra, note sotto ogni slide) ----------
os.makedirs(f'{out}/html', exist_ok=True); os.makedirs(f'{out}/pdf', exist_ok=True); os.makedirs(f'{out}/pptx', exist_ok=True)
body = ''
for k, (s, n) in enumerate(zip(slides, notes), 1):
    nt = f'<details class="nt"><summary>Note per l\'istruttore</summary><p>{H.escape(n)}</p></details>' if n else ''
    body += f'<div class="w" id="s{k}"><div class="f">{s}</div></div>{nt}'
view = f'''<!doctype html><html lang="it"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{H.escape(title)}</title>{links}<style>{BASE}
body{{background:#1B2A41;font-family:'Nunito Sans',Arial,sans-serif}} header{{color:#FFF8EE;padding:18px 16px;font-size:20px;font-weight:800;max-width:1200px;margin:0 auto}}
.w{{position:relative;max-width:1200px;margin:0 auto 10px;aspect-ratio:16/9;overflow:hidden;border-radius:10px;background:#FFF8EE}}
.f{{position:absolute;left:0;top:0;width:1920px;height:1080px;transform-origin:0 0}}
.nt{{max-width:1200px;margin:0 auto 28px;color:#C9D6E3;font-size:15px;line-height:1.5;padding:0 6px}} .nt summary{{cursor:pointer;color:#FFC145;font-weight:700}}
@media print{{body{{background:#fff}} header,.nt{{display:none}} .w{{break-after:page;border-radius:0}}}}</style></head>
<body><header>{H.escape(title)} · {len(slides)} slide</header>{body}
<script>function fit(){{document.querySelectorAll('.w').forEach(w=>{{w.firstChild.style.transform='scale('+(w.clientWidth/1920)+')'}})}}addEventListener('resize',fit);fit();</script></body></html>'''
open(f'{out}/html/{name}.html', 'w').write(view)

# ---------- pagina di stampa: slide a grandezza naturale ----------
printp = f'<!doctype html><html><head><meta charset="utf-8">{links}<style>{BASE} @page{{size:1920px 1080px;margin:0}} body{{margin:0}} section.s{{break-after:page}}</style></head><body>{"".join(slides)}</body></html>'
tmp = os.path.abspath(f'{deck}/_print.html'); open(tmp, 'w').write(printp)
prs = Presentation(); prs.slide_width = Emu(12192000); prs.slide_height = Emu(6858000)
with sync_playwright() as p:
    b = p.chromium.launch(executable_path=CHROME, args=['--no-sandbox','--ignore-certificate-errors'], proxy={'server': os.environ['HTTPS_PROXY']} if os.environ.get('HTTPS_PROXY') else None)
    pg = b.new_page(viewport={'width': 1920, 'height': 1080})
    pg.goto('file://' + tmp, wait_until='networkidle'); pg.evaluate('document.fonts.ready'); pg.wait_for_timeout(500)
    pg.pdf(path=f'{out}/pdf/{name}.pdf', width='1920px', height='1080px', print_background=True)
    secs = pg.query_selector_all('section.s')
    for k, s in enumerate(secs):
        png = s.screenshot(type='jpeg', quality=88)
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        sl.shapes.add_picture(io.BytesIO(png), 0, 0, prs.slide_width, prs.slide_height)
        if notes[k]: sl.notes_slide.notes_text_frame.text = notes[k]
    b.close()
prs.save(f'{out}/pptx/{name}.pptx'); os.remove(tmp)
print(name, len(slides), 'slide')
